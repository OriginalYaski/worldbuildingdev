#This window is for the purpose of setting up the population demographics of the campaign continent

import tkinter as tk
from openpyxl.styles import Border, Side
from openpyxl import Workbook, cell
from Universal_GUI_styling import *
from ScalingGUI import PageTwo

class PageOne(tk.Frame):

    #This function saves the generated values to the appropriate spot in the excel file and moves to the next page
    def submit(self, controller):
        strings = ["Active PC population", "Active NPC population", "Total active population"]

        for i in range(3):
            controller.ws['A' + str(i+1)].value = strings[i]
            controller.ws['A' + str(i+1)].border = LEFT_CELL
            controller.ws['B' + str(i+1)].border = RIGHT_CELL
        
        controller.ws['B3'].value = calcvars[3]
        controller.ws['B1'].value = calcvars[4]
        controller.ws['B2'].value = calcvars[3] - calcvars[4]
        
        controller.ws.column_dimensions['A'].width = 21
        controller.ws.column_dimensions['B'].width = 13
        
        controller.save()
        controller.frames[PageTwo].startup(controller)
        controller.show_frame(PageTwo)

    #This function updates the population counts for the total working population and the total working PC population
    def total(self):
        #Find first the total working population as a fraction of total population, then use the ratio of PCs to find the working PC population
        calcvars[3] = round(calcvars[0] * (calcvars[1]/100))
        calcvars[4] = round(calcvars[3] * (calcvars[2]/100))

        #Store both these values for display
        textvars[3].set(calcvars[3])
        textvars[4].set(calcvars[4])

    #This function stores the correct input variable, then updates the displays
    def update(self, target):

        #If the entry is for the total population, display it as is
        if target == 0:
            textvars[target].set(newvars[target].get())
        #Otherwise, display it as a percentage
        else:
            textvars[target].set(newvars[target].get() + "%")

        #Convert and store the variable as a float as well for further calculations
        calcvars[target] = float(newvars[target].get())

        #Update the calculations
        self.total()

        #Clear the submitted field
        newvars[target].set("")
        
    def __init__(self, parent, controller, defaults):
        #Initiate the master frame
        tk.Frame.__init__(self, parent)

        #Create the frames for the window
        topFrame = tk.Frame(self)
        middleFrame = tk.Frame(self)
        bottomFrame = tk.Frame(self, bd=3,
                               highlightbackground='black', highlightthickness=2)
        #Place the frames within the window
        topFrame.pack(side="top")
        middleFrame.pack(side="top")
        bottomFrame.pack(side="bottom")
        #Create more frames to hold the submission options
        popFrame = tk.Frame(topFrame, bd=3,
                            highlightbackground='black', highlightthickness=2)
        workFrame = tk.Frame(middleFrame, bd=3,
                             highlightbackground='black', highlightthickness=2)
        PCFrame = tk.Frame(middleFrame, bd=3,
                              highlightbackground='black', highlightthickness=2)
        leftFrame = tk.Frame(bottomFrame)
        rightFrame = tk.Frame(bottomFrame)
        #Place the subframes within the frames
        popFrame.pack(side="bottom")
        workFrame.pack(side="left")
        PCFrame.pack(side="right")
        leftFrame.pack(side="left")
        rightFrame.pack(side="right")

        #Initiate the lists that will be used to populate the window Labels
        global textvars, newvars, calcvars
        textvars = [tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar()]
        calcvars = [0,0,0,0,0]

        #Set the Labels based on the default values
        for i in range(3):
            textvars[i].set(defaults[i+1])
            calcvars[i] = float(defaults[i+1])

        #Convert the ratios into percentages
        textvars[1].set(textvars[1].get() + "%")
        textvars[2].set(textvars[2].get() + "%")

        #Run the calculations to find the total population values
        self.total()

        #Create fields to store user inputs
        newvars = [tk.StringVar(), tk.StringVar(), tk.StringVar()]

        #Title the window to give direction
        title = tk.Label(topFrame, text="Now, determine this world/continent's demographics.", font=LARGE_FONT).pack(side="top", pady=10, padx=10)

        #Set up the first frame to display the total population
        popGo = tk.Button(popFrame, text="Update population", command=lambda: self.update(0)).pack(side="bottom")
        popSet = tk.Entry(popFrame, textvariable = newvars[0])
        popSet.pack(side="bottom")
        pop = tk.Label(popFrame, textvariable=textvars[0]).pack(side="bottom")
        popTitle = tk.Label(popFrame, text="What is the total population?").pack(side="bottom")

        #Set up the next frame to display what percentage of the population is of working age
        workGo = tk.Button(workFrame, text="Update working class", command=lambda: self.update(1)).pack(side="bottom")
        workSet = tk.Entry(workFrame, textvariable = newvars[1])
        workSet.pack(side="bottom")
        working = tk.Label(workFrame, textvariable=textvars[1]).pack(side="bottom")
        workTitle = tk.Label(workFrame, text="What percentage of the population is of working age?").pack(side="bottom")

        #Set up the next frame to display how common PCs are in this world
        PCGo = tk.Button(PCFrame, text="Update character class", command=lambda: self.update(2)).pack(side="bottom")
        PCSet = tk.Entry(PCFrame, textvariable = newvars[2])
        PCSet.pack(side="bottom")
        PCs = tk.Label(PCFrame, textvariable=textvars[2]).pack(side="bottom")
        PCTitle = tk.Label(PCFrame, text="What percentage of the population Has PC levels?").pack(side="bottom")

        #Set up the final frame to display the final values
        submitButton = tk.Button(bottomFrame, text="Submit", command=lambda: self.submit(controller)).pack(side="bottom")
        submitLabel = tk.Label(bottomFrame, text="Does this look correct?", font=LARGE_FONT).pack(side="top")
        workName = tk.Label(leftFrame, text="Total population of working age.").pack(side="top")
        workTotal = tk.Label(leftFrame, textvariable=textvars[3]).pack(side="bottom")
        PCName = tk.Label(rightFrame, text="Total active population with PC levels.").pack(side="top")
        PCTotal = tk.Label(rightFrame, textvariable=textvars[4]).pack(side="bottom")
        
