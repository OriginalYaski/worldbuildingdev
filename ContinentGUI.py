#Top level GUI caller for the project

import tkinter as tk
from tkinter import font
from openpyxl import Workbook
from openpyxl import load_workbook
from DemographicsGUI import PageOne
from ScalingGUI import PageTwo
from ClasslistGUI import PageThree
from ClassgenGUI import PageFour

LARGE_FONT= ("Verdana", 12)
        
#------------------------------------------------------------------------

class MainWindow(tk.Tk):
        
#------------------------------------------------------------------------

    def __init__(self, *args, **kwargs):

        #Load the default population values
        with open("default values.txt") as f:
            self.defaultlist = f.read().splitlines()
        #Instantiate a list to hold the values we save/load from a separate file
        self.load_list = ["ND"] * 7
            
        global campaign, wb, ws
        campaign = ""
        
        tk.Tk.__init__(self, *args, **kwargs)
        self.title("TTRPG population generator")
        container = tk.Frame(self)

        container.pack(side="top", fill="both", expand = True)

        container.grid_rowconfigure(0, weight=1)
        container.grid_columnconfigure(0, weight=1)

        self.frames = []

        for F in (StartPage, PageOne, PageTwo, PageThree, PageFour):

            frame = F(container, self, self.defaultlist)

            self.frames.append(frame)

            frame.grid(row=0, column=0, sticky="nsew")

        self.show_frame(0)

        self.geometry("600x444")
        
#------------------------------------------------------------------------

    def show_frame(self, cont):
        '''if cont == PageThree:
            cont.setup(self.frames[cont], self)'''

        frame = self.frames[cont]
        frame.tkraise()
        
#------------------------------------------------------------------------

    def save(self):
        self.wb.save(self.campaign + '.xlsx')
        with open(self.campaign + '.txt', 'w') as f:
            for line in self.load_list:
                f.write(f"{line}\n")
        
#------------------------------------------------------------------------

    def finish(self):
        self.save()
        self.wb.close()

        root.destroy()
        
#------------------------------------------------------------------------

class StartPage(tk.Frame):
        
#------------------------------------------------------------------------

    def __init__(self, parent, controller, defaults):
        #default = tk.font.nametofont("TkDefaultFont")
        #print(default.actual())
        global titleFrame, messageVar, load_list
        tk.Frame.__init__(self,parent)
        #Create the frames for the window
        titleFrame = tk.Frame(self)
        enterFrame = tk.Frame(self)
        #Place the frames within the window
        titleFrame.pack(side="top")
        enterFrame.pack(side="bottom")

        updateVar = tk.StringVar()
        messageVar = tk.StringVar()

        title = tk.Label(titleFrame, text="Welcome to your new campaign! Please, give it a name.", font=LARGE_FONT).pack(side="top", pady=100)
        error_msg = tk.Message(titleFrame, textvariable = messageVar, width=500).pack(side="bottom")
        campaign_name = tk.Entry(enterFrame, textvariable = updateVar)
        campaign_name.pack(side="top")
        load_button = tk.Button(enterFrame, text="Load this Campaign", command=lambda: self.load(updateVar.get(), controller)).pack(side="left", padx=25, pady=(0,50))
        new_button = tk.Button(enterFrame, text="Start new Campaign", command=lambda: self.submit(updateVar.get(), controller)).pack(side="right", padx=25, pady=(0,50))
        
#------------------------------------------------------------------------

    def submit(self, title, controller):
        controller.campaign = title
        controller.wb = Workbook()
        controller.ws = controller.wb.active
        controller.ws.title = 'Class Array'
        controller.ws['A36'].value = 'Modification array:'
        for i in range(1,6):
            controller.ws.cell(row = 37, column = i, value = 0)
        controller.wb.create_sheet('Full Array')

        controller.save()
        controller.frames[1].startup(controller)
        controller.show_frame(1)
        #controller.frames[4].setup(controller)
        #controller.show_frame(3)
         
#------------------------------------------------------------------------

    def load(self, title, controller):

        controller.campaign = title
        #Try to load the Campaign, and display an error if unsuccessful
        try:
            #Load in the user friendly sheet
            controller.wb = load_workbook(title + '.xlsx')

            #Load in the externally stored values
            with open(title + '.txt') as f:
                controller.load_list = f.read().splitlines()
        except:
            messageVar.set("I couldn't find a Campaign of this name. Try doublechecking your spelling and confirming that the Campaign is saved in the correct directory.")
            return

        #If the campaign is found, move the active sheet to the sheet containing our progress tracker
        try:
            controller.ws = controller.wb['Class Array']
        except:
        #If the main sheet isn't found, you're probably not loading an actual campaign
            messageVar.set("This file doesn't seem to be a Campaign file. Please check your spelling to confirm that this is an actual Campaign.")
            return

        #First, load in all pages that have been visited
        for i in range(5,0,-1):
            if controller.ws.cell(row=37, column=i).value != 0:
                controller.frames[i].load_in(controller)
        
        #Find what page was most recently updated, but whose updates have not yet been passed to another page, and move the window to that page
        for i in range(1,6):
            if int(controller.ws.cell(row=37, column=i).value) == 2:
                controller.show_frame(i)
                return

        #If no page shows as updated, this sheet isn't working properly. Let the user know
        messageVar.set("This file doesn't seem to be formatted properly. Check the integrity of this file, or load another Campaign.")
        
#------------------------------------------------------------------------


'''with open("default values.txt") as f:
    defaultlist = f.read().splitlines()
    classList = defaultlist[0].split(", ")
    pop = int(defaultlist[1])
    ratio = int(defaultlist[2])'''

'''#Create the window
root = tk.Tk()
root.title("New campaign")
#Create the frames for the window
titleFrame = tk.Frame(root)
enterFrame = tk.Frame(root)
#Place the frames within the window
titleFrame.pack(side="top")
enterFrame.pack(side="bottom")

updateVar = tk.StringVar()

title = tk.Label(titleFrame, text="Welcome to your new campaign! Please, give it a name.").pack(side="top")
newCampaign = tk.Entry(enterFrame, textvariable = updateVar)
newCampaign.pack(side="left")
submitButton = tk.Button(enterFrame, text="Submit", command=submit).pack(side="right")
'''
global root
root = MainWindow()
root.title("New Campaign")
root.resizable(False, False)
root.mainloop()


