#This window is for the purpose of setting up how many individuals there are of each character level in this setting

import tkinter as tk
import random
from openpyxl import Workbook
from Universal_GUI_styling import *


class PageTwo(tk.Frame):

#------------------------------------------------------------------------

    #This function saves the generated values to the appropriate spot in the excel file and moves to the next page
    def submit(self, controller):

        #Write the values to the sheet
        for i in range(20):
            #Write value to cell
            controller.ws.cell(row = i+1, column = 3,
                                value = int(adultList[i].get()))
            controller.ws.cell(row = i+1, column = 4,
                                value = int(PCList[i].get()))        
        
        controller.save()
        controller.frames[3].startup(controller, controller.defaultlist[0].split(", "))
        controller.show_frame(3)

#------------------------------------------------------------------------

    #method to setup the PC and total per-level values
    def total_lists(self, sheet, total, display, falloff, scale):
        #The function used to generate the per-level values technically extends up to level infinity
        #This helps take all those people that should exist at levels 21+ and rounds them into the level 20 values
        tracker = total

        for i in range(0, 19):
            #Get value for the level
            val = round(total/(falloff**(i/scale))-total/(falloff**((i+1)/scale)))

#With certain sets of values, the per-level tracking will drop to below 0.5 while there is still individuals left to assign
#In this case, it is more satisfying to simply assign the remaining member, rather than having a random level 20 character outlier
            if tracker == 0:
                val = 0
            elif val == 0 and tracker > 0:
                val = tracker
            elif tracker - val < 0:
                val = val - tracker
                
            #Add value to list
            display[19-i].set(val)

            #print(val)

            #Increment tracker
            tracker -= val

        #We aren't going to epic levels, so keep all the rest
        display[0].set(tracker)

#------------------------------------------------------------------------

    def update(self, controller):

        #Setup the PC and total level distributions
        #The total_lists method has been set up to accept variable scaling, but this method isn't set up for that just yet
        self.total_lists(controller.ws, PC_total, PCList,
                         int(scalar.get()), int(frequency.get()))
        self.total_lists(controller.ws, Adult_total, adultList,
                         int(scalar.get()), int(frequency.get()))

        #Store the necessary values in the secondary file
        controller.load_list[2] = scalar.get()
        controller.load_list[3] = frequency.get()
        
        #Let the next page know that there has been an update
        controller.ws['B37'].value = 2
        
#------------------------------------------------------------------------

    def load_in(self, controller):
        #If we are loading in a campaign, retreive the lists stored in the master file
        for i in range(20):
            #Store values in the lists
            adultList[i].set(str(controller.ws.cell(row = i+1, column = 3).value))
            PCList[i].set(controller.ws.cell(row = i+1, column = 4).value)

        #Get the population totals as well
        PC_total = int(controller.ws.cell(row = 1, column = 2).value)
        Adult_total = int(controller.ws.cell(row = 3, column = 2).value)
        #Get the scalar and frequency values from the secondary file
        scalar.set(controller.load_list[2])
        frequency.set(controller.load_list[3])


#------------------------------------------------------------------------

    def startup(self, controller):

        global PC_total, Adult_total

        #Run update if a change was made on the previous page
        if controller.ws['A37'].value == 2:

            if controller.ws['B37'].value == 0:
            #If this is first time starting up this page, perform initial sheet formatting
                controller.ws['C21'].value = "Adults per level"
                controller.ws['C21'].border = BOTTOM_CELL
                controller.ws.column_dimensions['C'].width = 15
                
                controller.ws['D21'].value = "PCs per level"
                controller.ws['D21'].border = BOTTOM_CELL
                controller.ws.column_dimensions['D'].width = 13                    

                # Apply style
                for i, rowOfCellObjects in enumerate(controller.ws['C2':'D20']):
                    for n, cellObj in enumerate(rowOfCellObjects):
                        cellObj.border = COLUMN
                controller.ws['C1'].border = TOP_CELL
                controller.ws['D1'].border = TOP_CELL

            #Get the totals when the page loads
            PC_total = int(controller.ws.cell(row = 1, column = 2).value)
            Adult_total = int(controller.ws.cell(row = 3, column = 2).value)

            #Update the displayed lists
            self.update(controller)

            #Mark that the contents of the previous page have been received
            controller.ws['A37'].value = 1
        
#------------------------------------------------------------------------

    def __init__(self, parent, controller, defaults):
        tk.Frame.__init__(self, parent)

        #Declare the global variables that will be needed elsewhere
        global scalar, frequency, adultList, PCList

        #Instantiate the variables we will set on startup
        PC_total = 1
        Adult_total = 1

        #Instantiate the variables we will use for calculating the lists, and set them to the default values
        scalar = tk.StringVar()
        scalar.set(defaults[4])
        frequency = tk.StringVar()
        frequency.set(defaults[5])

        #Instantiate the lists where we will store the per-level lists
        adultList = []
        PCList = []
        #Instantiate the matrix that will store the displayed output windows
        displayMatrix = []
        n = 0
        #Fill the lists with variables, create the display windows, and place them in the frame
        for i in range(20):
            if i == 11:
                n = 4
            adultList.append(tk.StringVar())
            PCList.append(tk.StringVar())

            row = []
            row.append(tk.Entry(self, textvariable=adultList[i],
                                width=10).place(y=40+i*20,x=432))
            row.append(tk.Label(self, text=20-i).place(
                y=38+i*20, x=504+n))
            row.append(tk.Entry(self, textvariable=PCList[i],
                                width=10).place(y=40+i*20,x=531))
            displayMatrix.append(row)

        window_label = tk.Label(self, text="Determine the scaling of your setting",
                               font=LARGE_FONT).place(x=60, y=15)
        explanation = tk.Message(self, text="In most settings, higher level characters are rarer than lower level characters. "
"This ensures that the higher in level something is, the rarer it is, and thus the more special it feels for players. "
"This section determines how many characters of each level exist in this setting.\n\n"
"This page sets this up through a simple exponential curve. "
"The 'Factor' number on the left determines the falloff, and the 'Frequency' number on the right determines how often this is applied. "
"As an example, if the number on the left is 2, and the number on the right is 1, then each level will have half the population of the level before it (1/2 every 1 level). "
"If the Factor is 5, and the Frequency is 3, then each level's population will be one-fifth of the population three levels before it (Level 1 will have 5-times the population of level 4, 25-times the population of level 7, etc.).\n\n"
"Also, if you want, you can manually set any of the fields along the right side of this window, if you want to personally tweak your level distribution "
"(WARNING: If you add population to any level, make sure you subtract an equal amount from another level, and vice-versa. Failure to do so may cause errors).",
                               width=380).place(x=20, y=40)
        first_col_label1 = tk.Label(self, text = "Adults per level").place(
            y=4, x=422)
        first_col_label2 =tk.Label(self, text = "(PC AND NPC)",
                                   font=("Segoe UI", 7)).place(
            y=22, x=433) 
        second_col_label = tk.Label(self, text = "PCs per level").place(
            y=10, x=525)
        scalar_label = tk.Label(self, text = "Factor:").place(
            y=360, x=90)
        scalar_field = tk.Entry(self, textvariable=scalar, width=5).place(
            y=361, x=130)
        freq_label = tk.Label(self, text = "Frequency:").place(
            y=360, x=190)
        freq_field = tk.Entry(self, textvariable=frequency, width=3).place(
            y=361, x=252)
        updateButton = tk.Button(self, text="Calculate",
                                 command=lambda: self.update(controller)).place(
                                     y=357,
                                     x=300)
        submitButton = tk.Button(self, text="Back",
                                 command=lambda: controller.show_frame(1)).place(
                                     y=400,
                                     x=100)
        submitButton = tk.Button(self, text="Submit",
                                 command=lambda: self.submit(controller)).place(
                                     y=400,
                                     x=300)
        
