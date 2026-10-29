#This window is for the purpose of setting up the available classes in this setting

import tkinter as tk
from openpyxl.styles import Border, Side
from openpyxl import Workbook
from Universal_GUI_styling import *

class PageThree(tk.Frame):
        
#------------------------------------------------------------------------
    
    #Checks if any changes have been made and if we are clear to proceed to the next page
    def submit(self, controller):
        #Check whether anything has been overwritten and if are clear to proceed
        check_conditions = self.check_updated(controller)
                        
        #Only proceed if at least 1 class has been selected
        if check_conditions:

            #If at least 1 class has been selected, save the current sheet and proceed
            controller.save()
            controller.frames[4].startup(controller)
            controller.show_frame(4)

#------------------------------------------------------------------------

    #This function writes the class list to the excel sheet
    def write_list(self, controller, classes):
        
        #Write the new class list to the sheet
        for i in range(len(classes)):
            controller.ws.cell(row = 21, column = i+6, value = classes[i])

        #Store the ENTIRE list (even those not selected) to the secondary file
        second_list = classList[0]
        for i in range(1, len(classList)):
            second_list = second_list + ", " + classList[i]
        controller.load_list[4] = second_list

        #Store the length of the new class list and mark that it has been updated
        controller.ws['A6'].value = len(classes)
        controller.ws['C37'].value = 2

#------------------------------------------------------------------------

    def check_updated(self, controller):

        #Check which classes have been selected, and store them in a separate list
        writeList = []
        for i in range(len(varList)):
            if varList[i].get() == 1:
                writeList.append(classList[i])
                
        #If no classes have been selected, return that information
        if writeList == []:
            return False
        
        else:
            
        #If there is at least one class selected, and this is our first time passing through, we can proceed automatically
            if controller.ws['C37'].value == 0.5:
                #Perform setup stylization to the sheet
                controller.ws['A5'].value = "Number of classes"
                controller.ws['A5'].border = TOP_CELL
                controller.ws['A6'].border = BOTTOM_CELL

                #Write the list to the sheet
                self.write_list(controller, writeList)
                
                
        #If there is a list already written to the file, check if the lengths match
            elif len(writeList) == controller.ws['A6'].value:
                #If the lengths match, check if there are any discrepancies between the new list and the old list
                for i in range(len(writeList)):
                    
                    #If there are any discrepancies, overwrite the list
                    if controller.ws.cell(row = 21, column = i+6).value not in writeList:

                        #Write the new list (we know the lengths match, so we can just overwrite it)
                        self.write_list(controller, writeList)

                        #Then exit the loop
                        break

        #If the lengths don't match, then we know that the list has been overwritten
            else:
                
                #First, clear out the old list
                for i in range(controller.ws['A6'].value):
                    controller.ws.cell(row = 21, column = i+6).value = None

                #Then write the new list
                self.write_list(controller, writeList)
            
        #Whether or not an update happened, confirm that we are clear to pass on to the next page
            return True
        
#------------------------------------------------------------------------

    #When the button is pressed, update the displayed and stored lists
    def update_list(self, frame):
        global r, c
        #Create a position in both lists, one to save the class name, and one to save the checkbox current state
        classList.append(updateVar.get())
        varList.append(tk.IntVar())

        #Create a checkbox and place it with the new class name
        tk.Checkbutton(frame, text=classList[-1], variable=varList[-1]).grid(row=r, column=c)
        c += 1
        if c > 4:
            r += 1
            c = 0

        #To-do: Implement option to save class to the default list (to auto-load it into future campaigns)

        updateVar.set("")
        
#------------------------------------------------------------------------

    def select_all(self):
        for num in varList:
            num.set(1)
        
#------------------------------------------------------------------------

    def select_none(self):
        for num in varList:
            num.set(0)
        
#------------------------------------------------------------------------

    def load_in(self, controller):
        #If we are loading in a campaign, load in all the classes in the previous list (even those that weren't selected)
        new_list = controller.load_list[4].split(", ")
        self.startup(controller, new_list, True)

        #Find which ones were selected, and set those boxes to checked
        for i in range(controller.ws['A6'].value):
            classname = controller.ws.cell(row = 21, column = i+6).value
        #Every class name that was written to the excel file SHOULD be present in the secondary file
        #If it's not, then the users were meddling with one of the two files, and I at this time I am not about to correct their mistake (maybe I'll come back to it once this project is further along)
            if classname in classList:
                varList[classList.index(classname)].set(1)
        
#------------------------------------------------------------------------

    def startup(self, controller, defaults, load=False):
        global listFrame
        #Load in each class name in the list, and update the variables and display for each name
        if load or controller.ws['C37'].value == 0:
            for classname in defaults:
                updateVar.set(classname)
                self.update_list(listFrame)

        if not load:
        #Set page state to special passthrough variable (page has loaded, but has not yet reached its first commit)
            controller.ws['C37'].value = 0.5
        
#------------------------------------------------------------------------

    def __init__(self, parent, controller, defaults):
        tk.Frame.__init__(self, parent)
        global classList, varList, checkList, updateVar, c, r, listFrame
        c = r = pointer = 0
        classList = []
        varList = []
        
        #Create the frames for the window
        listFrame = tk.Frame(self)
        bottomFrame = tk.Frame(self)
        #Place the frames within the window
        listFrame.pack(side="top")
        bottomFrame.pack(side="bottom")
        #Place this button to the far left before creating more frames
        back_button = tk.Button(bottomFrame, text="Back", command=lambda: controller.show_frame(2)).pack(side="left")
        #Create more frames to place within the bottom frame
        entryFrame = tk.Frame(bottomFrame)
        submitFrame = tk.Frame(bottomFrame)
        selectFrame = tk.Frame(submitFrame)
        #Place the frames within the bottom frame
        entryFrame.pack(side="left")
        submitFrame.pack(side="right")
        selectFrame.pack(side="left")
                
        #The variable where the new class name is stored
        updateVar = tk.StringVar()

        #Create and display the elements to add a new class to the list, and to submit the final list
        newLabel = tk.Label(entryFrame, text="Add another class to the list").pack(side="top")
        newClass = tk.Entry(entryFrame, textvariable = updateVar)
        newClass.pack(side="left")
        classButton = tk.Button(entryFrame, text="Add", command=lambda: self.update_list(listFrame)).pack(side="left")
        submitButton = tk.Button(submitFrame, text="Submit", command=lambda: self.submit(controller)).pack(side="right")
        allButton = tk.Button(selectFrame, text="Select All", command=self.select_all).pack(side="top")
        noneButton = tk.Button(selectFrame, text="Select None", command=self.select_none).pack(side="bottom", padx=20)
